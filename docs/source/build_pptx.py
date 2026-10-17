import sys, os, re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.dml import MSO_LINE

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import content as T
DECK = HERE + "/assets"; OUT = sys.argv[1]
def C(h): return RGBColor.from_string(h.lstrip("#"))
INK, DEEP, TEAL, GREY, RULE, WHITE = (C(x) for x in ("111A20", "00526F", "5390A2", "5F6F78", "D3DDE2", "FFFFFF"))
HEAD, BODY = "Athelas", "Frutiger"
LOGO_W, LOGO_H = 0.32 * 2000 / 276, 0.32          # logo is 2000x276
L, R = 1.0, 12.333                                # content margins
W = R - L
FS, MIN_PT = 0.84, 10          # every size below is scaled by FS, for more white space
from PIL import ImageFont
_F = "/Library/Fonts/Frutiger.ttf"
def width_in(t, pt):           # measured width of a Frutiger line, in inches
    return ImageFont.truetype(_F, int(pt * 10)).getlength(t) / 10 / 72

prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
LAYOUT = prs.slide_layouts[6]

def new_slide(notes):
    s = prs.slides.add_slide(LAYOUT)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    if not os.environ.get("NO_NOTES"):
        s.notes_slide.notes_text_frame.text = notes
    return s

def text(s, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP):
    b = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = b.text_frame; tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", PP_ALIGN.LEFT)
        if "after" in p: para.space_after = Pt(p["after"])
        if "line" in p: para.line_spacing = p["line"]
        runs = p["runs"] if "runs" in p else [(p["text"], {})]
        runs = [(piece, dict(o, **({"italic": True} if piece == "Noéa" else {})))     # the name is always italic
                for t, o in runs for piece in re.split(r"(Noéa)", t) if piece]
        for t, o in runs:
            r = para.add_run(); r.text = t
            r.font.name = o.get("font", p.get("font", BODY))
            r.font.size = Pt(max(MIN_PT, round(o.get("size", p.get("size", 18)) * FS, 1)))
            r.font.bold = o.get("bold", p.get("bold", False))
            r.font.italic = o.get("italic", p.get("italic", False))
            r.font.color.rgb = o.get("color", p.get("color", INK))
            spc = o.get("spc", p.get("spc"))
            if spc: r._r.get_or_add_rPr().set("spc", str(spc))
    return b

def hline(s, x, y, w, color=RULE, weight=1.0, dash=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y))
    c.line.color.rgb = color; c.line.width = Pt(weight)
    if dash: c.line.dash_style = MSO_LINE.DASH

def vline(s, x, y, h, color=RULE, weight=1.0):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x), Inches(y + h))
    c.line.color.rgb = color; c.line.width = Pt(weight)

def subtitle(s, y, t, base=16, color=GREY):
    """One line, always: shrink until the measured width fits inside the content width."""
    pt = base
    while width_in(t, pt * FS) > W * 0.97 and pt > 11: pt -= 0.5
    assert width_in(t, pt * FS) <= W, (t, pt)
    print(f"  subtitle {pt*FS:.1f}pt, {width_in(t, pt*FS):.2f}in of {W:.2f}in: {t[:48]}...")
    text(s, L, y, W, 0.4, [dict(text=t, size=pt, color=color)])

def label(s, x, y, w, t, color=TEAL):
    text(s, x, y, w, 0.3, [dict(text=t.upper(), size=12, bold=True, color=color, spc=140)])

def title(s, t, size=40, y=0.75):
    text(s, L, y, W, 1.5, [dict(text=t, size=size, font=HEAD, line=1.05)])

def footer(s, n=None):
    s.shapes.add_picture(f"{DECK}/logo.png", Inches(L), Inches(6.82), height=Inches(LOGO_H))
    if n: text(s, R - 0.7, 6.9, 0.7, 0.3, [dict(text=str(n), size=11, color=GREY, align=PP_ALIGN.RIGHT)])

# ---------------- 1. cover ----------------
s = new_slide(T.SPEAKER_NOTES[1])
s.shapes.add_picture(f"{DECK}/noea_wordmark.png", Inches(L), Inches(0.85), height=Inches(1.15))
text(s, L, 2.4, W, 0.7, [dict(text=T.COVER_TAGLINE, size=30, font=HEAD, italic=True, color=DEEP)])
text(s, L, 3.05, W, 0.4, [dict(text=T.COVER_SUBTITLE, size=16, color=GREY)])
ECG_IDS = []
if not os.environ.get("NO_ECG"):          # NO_ECG=1 builds the cover without the strip (base frame for the video)
    pic = s.shapes.add_picture(f"{DECK}/ecg_row1.png", Inches(L), Inches(4.2), width=Inches(W))
    ECG_IDS = [pic.shape_id]
text(s, L, 6.45, W, 0.3, [dict(text=T.COVER_CAPTION, size=11, color=GREY)])
s.shapes.add_picture(f"{DECK}/logo.png", Inches(R - LOGO_W), Inches(0.95), height=Inches(LOGO_H))
COVER = s

# ---------------- 2. the gap ----------------
s = new_slide(T.SPEAKER_NOTES[2])
s.shapes.add_picture(f"{DECK}/{T.GAP_DECOR}", Inches(R + 1.0 - 5.13), 0, width=Inches(5.13), height=Inches(7.5))   # decorative panel on the right edge, behind the text
text(s, L, 2.2, W, 2.4, [dict(text=t, size=46, font=HEAD, line=1.1) for t in T.GAP_TITLE])
hline(s, L, 4.6, 1.2, DEEP, 3)
text(s, L, 4.85, W, 1.0, [dict(text=T.GAP_LINE, size=22, font=HEAD, italic=True, color=DEEP, line=1.15)])
footer(s, 2)

# ---------------- 3. how it reads ----------------
s = new_slide(T.SPEAKER_NOTES[3])
title(s, T.READ_TITLE, size=36)
cw, gap, y0 = 2.6, 0.31, 2.75
for i, (n, h, paras, hot) in enumerate(T.READ_STEPS):
    x = L + i * (cw + gap)
    hline(s, x, y0, cw, DEEP if hot else RULE, 3 if hot else 1.25)
    text(s, x, y0 + 0.25, cw, 0.3, [dict(text=n, size=13, bold=True, color=DEEP if hot else TEAL, spc=100)])
    text(s, x, y0 + 0.7, cw, 0.9, [dict(text=h, size=23, font=HEAD, line=1.0)])
    text(s, x, y0 + 1.55, cw - 0.05, 2.2, [dict(text=t, size=14.5, color=GREY if not hot else INK, line=1.12, after=7) for t in paras])
assert width_in(T.READ_SOURCE, 10 * FS) <= W, "source line too long for one line"
text(s, L, 6.3, W, 0.3, [dict(text=T.READ_SOURCE, size=10, color=GREY)])
footer(s, 3)

# ---------------- 4. universal ----------------
s = new_slide(T.SPEAKER_NOTES[4])
title(s, T.UNIV_TITLE)
text(s, L, 2.25, W, 1.5, [dict(text=T.UNIV_LINE, size=27, font=HEAD, line=1.15)])
y = 3.6
for name, detail, dim in T.UNIV_ROWS:
    hline(s, L, y, W, RULE, 1.25)
    text(s, L, y + 0.2, 4.2, 0.5, [dict(text=name, size=21, font=HEAD, color=GREY if dim else INK)])
    text(s, L + 4.4, y + 0.27, W - 4.4, 0.5, [dict(text=detail, size=14, color=GREY)])
    y += 0.8
footer(s, 4)

# ---------------- 5. two layers ----------------
s = new_slide(T.SPEAKER_NOTES[5])
title(s, T.LAYERS_TITLE)
colw = 5.2
def layer(x, head, tag, tagcolor, bullets):
    text(s, x, 2.55, colw, 0.6, [dict(text=head, size=32, font=HEAD)])
    label(s, x, 3.2, colw, tag, tagcolor)
    text(s, x, 3.8, colw, 2.8, [dict(text=b, size=16.5, color=INK, line=1.12, after=12) for b in bullets])
layer(L, *T.LAYER_SCIENCE[:2], GREY, T.LAYER_SCIENCE[2])
vline(s, 6.667, 2.6, 3.9)
layer(7.133, *T.LAYER_CURATORIAL[:2], DEEP, T.LAYER_CURATORIAL[2])
footer(s, 5)

# ---------------- 6. machine learning ----------------
s = new_slide(T.SPEAKER_NOTES[6])
title(s, T.ML_TITLE)
subtitle(s, 1.6, T.ML_SUBTITLE)
s.shapes.add_picture(f"{DECK}/anomaly_ribbon.png", Inches(L), Inches(2.35), width=Inches(W))
text(s, L, 5.6, W, 0.4, [dict(runs=[(T.ML_CALLOUT_BOLD + " ", dict(bold=True)), (T.ML_CALLOUT_REST, {})], size=15)])
for t in (T.ML_METHOD, T.ML_LIMIT): assert width_in(t, 10 * FS) <= W, ("one line only", t)
text(s, L, 6.1, W, 0.6, [dict(text=T.ML_METHOD, size=10, color=GREY, line=1.08, after=3),
                          dict(text=T.ML_LIMIT, size=10, color=GREY, line=1.08)])
footer(s, 6)

# ---------------- 7. the end ----------------
s = new_slide(T.SPEAKER_NOTES[7])
s.shapes.add_picture(f"{DECK}/{T.END_DECOR}", Inches(7.9), Inches(-0.6), width=Inches(9), height=Inches(9))   # decorative, behind the text; the rest of the slide stays plain white
title(s, T.END_TITLE)
TW = 8.6
text(s, L, 2.6, TW, 1.6, [dict(text=T.END_LINE, size=26, font=HEAD, line=1.15)])
hline(s, L, 4.15, 1.2, DEEP, 3)
text(s, L, 4.4, TW, 1.2, [dict(text=T.END_RESEARCH, size=22, font=HEAD, italic=True, color=DEEP, line=1.15)])
text(s, R - 4.5, 6.9, 3.7, 0.3, [dict(text=T.END_URL_TEXT, size=11, bold=True, color=DEEP, align=PP_ALIGN.RIGHT)])
footer(s, 7)

def animate_cover(slide, ids, dur=3200, start=500):
    from lxml import etree
    NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
    n = [2]
    def nid():
        n[0] += 1; return n[0]
    effects = ""
    for k, spid in enumerate(ids):
        node = "afterEffect" if k == 0 else "withEffect"
        delay = start + k * dur
        a, b, c = nid(), nid(), nid()
        effects += (f'<p:par><p:cTn id="{a}" presetID="22" presetClass="entr" presetSubtype="8" fill="hold" nodeType="{node}">'
                    f'<p:stCondLst><p:cond delay="{delay}"/></p:stCondLst><p:childTnLst>'
                    f'<p:set><p:cBhvr><p:cTn id="{b}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
                    f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
                    f'<p:to><p:strVal val="visible"/></p:to></p:set>'
                    f'<p:animEffect transition="in" filter="wipe(left)"><p:cBhvr><p:cTn id="{c}" dur="{dur}"/>'
                    f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>'
                    f'</p:childTnLst></p:cTn></p:par>')
    xml = (f'<p:timing xmlns:p="{NS}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
           f'<p:par><p:cTn id="{nid()}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>'
           f'<p:par><p:cTn id="{nid()}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
           f'{effects}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>'
           f'</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
           f'</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')
    slide._element.append(etree.fromstring(xml))

if os.environ.get("ANIMATE") and ECG_IDS:
    animate_cover(COVER, ECG_IDS, dur=7000)

prs.core_properties.title = "Noéa — an introduction"
prs.core_properties.author = "Meeting of Waters"
prs.save(OUT); print("saved", OUT)
